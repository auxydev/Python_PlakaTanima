"""Raporlar: saatlik/günlük trafik, gelir, araç tipi dağılımı ve dışa aktarma."""

from __future__ import annotations

import csv
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from sqlalchemy import select

from plakatanima.core.formatting import format_datetime, format_duration
from plakatanima.core.plate import format_plate
from plakatanima.core.vehicle import VehicleType
from plakatanima.db.database import Database
from plakatanima.db.models import ParkingSession, PaymentMethod, SessionStatus


@dataclass(slots=True)
class DayReport:
    day: date
    entries: int = 0
    exits: int = 0
    revenue: Decimal = Decimal(0)


@dataclass(slots=True)
class ReportSummary:
    start: date
    end: date  # dahil
    entries: int = 0
    exits: int = 0
    revenue: Decimal = Decimal(0)
    unpaid_amount: Decimal = Decimal(0)
    unpaid_count: int = 0
    unique_plates: int = 0
    average_stay: timedelta = timedelta(0)
    exit_missed: int = 0
    manual_entries: int = 0
    days: list[DayReport] = field(default_factory=list)
    hourly_entries: list[int] = field(default_factory=lambda: [0] * 24)
    by_type: dict[VehicleType, int] = field(default_factory=dict)
    revenue_by_method: dict[PaymentMethod, Decimal] = field(default_factory=dict)

    @property
    def day_count(self) -> int:
        return (self.end - self.start).days + 1

    @property
    def peak_hour(self) -> int | None:
        if not any(self.hourly_entries):
            return None
        return max(range(24), key=self.hourly_entries.__getitem__)

    @property
    def average_daily_revenue(self) -> Decimal:
        return (self.revenue / self.day_count).quantize(Decimal("0.01"))


def _bounds(start: date, end: date) -> tuple[datetime, datetime]:
    return datetime.combine(start, time.min), datetime.combine(end + timedelta(days=1), time.min)


class ReportService:
    def __init__(self, db: Database) -> None:
        self._db = db

    def hourly_traffic(self, day: date) -> tuple[list[int], list[int]]:
        """Belirtilen gün için saat başına (giriş, çıkış) sayıları."""
        lo, hi = _bounds(day, day)
        entries, exits = [0] * 24, [0] * 24
        with self._db.session() as s:
            for (entry_time,) in s.execute(
                select(ParkingSession.entry_time).where(
                    ParkingSession.entry_time >= lo, ParkingSession.entry_time < hi
                )
            ):
                entries[entry_time.hour] += 1
            for (exit_time,) in s.execute(
                select(ParkingSession.exit_time).where(
                    ParkingSession.status == SessionStatus.CLOSED,
                    ParkingSession.exit_time >= lo,
                    ParkingSession.exit_time < hi,
                )
            ):
                exits[exit_time.hour] += 1
        return entries, exits

    def summary(self, start: date, end: date) -> ReportSummary:
        """[start, end] gün aralığının özeti (yerel saat)."""
        if end < start:
            start, end = end, start
        lo, hi = _bounds(start, end)
        report = ReportSummary(start=start, end=end)
        days = {start + timedelta(days=i): DayReport(start + timedelta(days=i))
                for i in range((end - start).days + 1)}  # fmt: skip
        plates: set[str] = set()
        by_type: Counter[VehicleType] = Counter()
        stay_total = timedelta(0)
        stay_count = 0

        with self._db.session() as s:
            entered = s.scalars(
                select(ParkingSession).where(
                    ParkingSession.entry_time >= lo, ParkingSession.entry_time < hi
                )
            )
            for session in entered:
                report.entries += 1
                days[session.entry_time.date()].entries += 1
                report.hourly_entries[session.entry_time.hour] += 1
                by_type[session.vehicle_type] += 1
                plates.add(session.plate)
                report.manual_entries += session.is_manual_entry
                report.exit_missed += session.status is SessionStatus.EXIT_MISSED

            exited = s.scalars(
                select(ParkingSession).where(
                    ParkingSession.status == SessionStatus.CLOSED,
                    ParkingSession.exit_time >= lo,
                    ParkingSession.exit_time < hi,
                )
            )
            for session in exited:
                report.exits += 1
                days[session.exit_time.date()].exits += 1
                stay_total += session.exit_time - session.entry_time
                stay_count += 1
                if not session.is_paid and session.fee:
                    report.unpaid_count += 1
                    report.unpaid_amount += session.fee

            paid = s.scalars(
                select(ParkingSession).where(
                    ParkingSession.is_paid.is_(True),
                    ParkingSession.paid_at >= lo,
                    ParkingSession.paid_at < hi,
                )
            )
            methods: dict[PaymentMethod, Decimal] = {}
            for session in paid:
                fee = Decimal(session.fee or 0)
                report.revenue += fee
                days[session.paid_at.date()].revenue += fee
                if session.payment_method is not None:
                    methods[session.payment_method] = (
                        methods.get(session.payment_method, Decimal(0)) + fee
                    )

        report.days = list(days.values())
        report.unique_plates = len(plates)
        report.by_type = dict(by_type.most_common())
        report.revenue_by_method = dict(sorted(methods.items(), key=lambda kv: -kv[1]))
        if stay_count:
            report.average_stay = timedelta(seconds=stay_total.total_seconds() // stay_count)
        return report


# ------------------------------------------------------------------ dışa aktarma
SESSION_HEADERS = [
    "Kayıt No",
    "Plaka",
    "Araç Tipi",
    "Durum",
    "Giriş",
    "Çıkış",
    "Süre",
    "Süre (dk)",
    "Ücret",
    "Ödeme",
    "Ödendi",
    "Giriş Kamerası",
    "Çıkış Kamerası",
    "Sahibi",
    "Not",
]


def _session_row(session: ParkingSession) -> list:
    vehicle = session.vehicle
    duration = session.duration()
    return [
        session.id,
        format_plate(session.plate),
        session.vehicle_type.label,
        session.status.label,
        session.entry_time,
        session.exit_time,
        format_duration(duration),
        int(duration.total_seconds() // 60),
        float(session.fee) if session.fee is not None else None,
        session.payment_method.label if session.payment_method else "",
        "Evet" if session.is_paid else "Hayır",
        session.entry_camera + (" (manuel)" if session.is_manual_entry else ""),
        session.exit_camera + (" (manuel)" if session.is_manual_exit else ""),
        vehicle.owner_name if vehicle else "",
        session.note,
    ]


def export_sessions_csv(sessions: list[ParkingSession], path: Path) -> Path:
    """Excel'in Türkçe ayarlarıyla doğrudan açılabilen CSV (UTF-8 BOM, ``;`` ayraç)."""
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(SESSION_HEADERS)
        for session in sessions:
            row = _session_row(session)
            row[4] = format_datetime(row[4], with_seconds=True)
            row[5] = format_datetime(row[5], with_seconds=True) if row[5] else ""
            row[8] = "" if row[8] is None else f"{row[8]:.2f}".replace(".", ",")
            writer.writerow(row)
    return path


def export_sessions_xlsx(
    sessions: list[ParkingSession],
    path: Path,
    *,
    summary: ReportSummary | None = None,
    currency: str = "TL",
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Geçişler"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")

    def style_header(ws, widths: list[int]) -> None:
        for col, width in enumerate(widths, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    sheet.append(SESSION_HEADERS)
    money_format = f'#,##0.00 "{currency}"'
    for session in sessions:
        sheet.append(_session_row(session))
        row = sheet.max_row
        sheet.cell(row=row, column=5).number_format = "dd.mm.yyyy hh:mm:ss"
        sheet.cell(row=row, column=6).number_format = "dd.mm.yyyy hh:mm:ss"
        sheet.cell(row=row, column=9).number_format = money_format
    style_header(sheet, [9, 13, 16, 14, 19, 19, 14, 10, 12, 12, 8, 20, 20, 18, 24])

    if summary is not None:
        daily = workbook.create_sheet("Günlük Özet")
        daily.append(["Tarih", "Giriş", "Çıkış", "Tahsilat"])
        for day in summary.days:
            daily.append([day.day, day.entries, day.exits, float(day.revenue)])
            daily.cell(row=daily.max_row, column=1).number_format = "dd.mm.yyyy"
            daily.cell(row=daily.max_row, column=4).number_format = money_format
        style_header(daily, [14, 10, 10, 16])

        totals = workbook.create_sheet("Özet")
        rows = [
            ("Dönem", f"{summary.start:%d.%m.%Y} – {summary.end:%d.%m.%Y}"),
            ("Toplam giriş", summary.entries),
            ("Toplam çıkış", summary.exits),
            ("Tekil araç", summary.unique_plates),
            ("Tahsilat", float(summary.revenue)),
            ("Ödenmemiş tutar", float(summary.unpaid_amount)),
            ("Ortalama kalış", format_duration(summary.average_stay)),
        ]
        rows += [(f"Araç: {vtype.label}", count) for vtype, count in summary.by_type.items()]
        rows += [
            (f"Ödeme: {method.label}", float(amount))
            for method, amount in summary.revenue_by_method.items()
        ]
        totals.append(["Gösterge", "Değer"])
        for label, value in rows:
            totals.append([label, value])
            if isinstance(value, float):
                totals.cell(row=totals.max_row, column=2).number_format = money_format
        style_header(totals, [28, 26])

    workbook.save(path)
    return path


def sessions_for_export(db: Database, start: date, end: date) -> list[ParkingSession]:
    from sqlalchemy.orm import selectinload

    lo, hi = _bounds(start, end)
    with db.session() as s:
        stmt = (
            select(ParkingSession)
            .options(selectinload(ParkingSession.vehicle))
            .where(ParkingSession.entry_time >= lo, ParkingSession.entry_time < hi)
            .order_by(ParkingSession.entry_time)
        )
        return list(s.scalars(stmt))
