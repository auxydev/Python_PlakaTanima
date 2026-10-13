from __future__ import annotations

import csv
import random
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from plakatanima.config import AppConfig, CameraRole
from plakatanima.core.vehicle import VehicleType
from plakatanima.db import Database, PaymentMethod, UserRole
from plakatanima.demo.seed import seed_demo_data
from plakatanima.services.auth import (
    AuthError,
    UserService,
    hash_password,
    verify_password,
)
from plakatanima.services.maintenance import (
    BackupError,
    backup_database,
    backup_if_due,
    list_backups,
)
from plakatanima.services.parking import ParkingService, PlateObservation
from plakatanima.services.reports import (
    ReportService,
    export_sessions_csv,
    export_sessions_xlsx,
    sessions_for_export,
)

DAY = date(2026, 5, 4)


@pytest.fixture
def db():
    database = Database("sqlite://")
    database.create_schema()
    yield database
    database.dispose()


def _pass(service, plate, vtype, entry, minutes, method=PaymentMethod.CARD):
    service.process(PlateObservation(plate, "Giriş", CameraRole.ENTRY, vtype, timestamp=entry))
    exit_time = entry + timedelta(minutes=minutes)
    result = service.process(PlateObservation(plate, "Çıkış", CameraRole.EXIT, timestamp=exit_time))
    if result.needs_payment and method is not None:
        service.record_payment(result.session_id, method)
    return result


@pytest.fixture
def filled(db):
    now = datetime.combine(DAY, datetime.min.time()) + timedelta(hours=20)
    service = ParkingService(db, AppConfig(), clock=lambda: now)
    base = datetime.combine(DAY, datetime.min.time())
    _pass(service, "34A1111", VehicleType.CAR, base + timedelta(hours=8), 90)  # 90 TL
    _pass(service, "34B2222", VehicleType.CAR, base + timedelta(hours=8, minutes=30), 60,
          PaymentMethod.CASH)  # 60 TL  # fmt: skip
    _pass(service, "06C3333", VehicleType.TRUCK, base + timedelta(hours=17), 120)  # 225 TL
    _pass(service, "35D4444", VehicleType.MOTORCYCLE, base + timedelta(hours=17), 10)  # 0
    _pass(service, "16E5555", VehicleType.CAR, base + timedelta(hours=18), 61, None)  # ödenmedi
    service.process(PlateObservation("41F6666", timestamp=base + timedelta(hours=19)))
    return db


def test_hourly_traffic(filled):
    entries, exits = ReportService(filled).hourly_traffic(DAY)
    assert entries[8] == 2 and entries[17] == 2 and entries[18] == 1 and entries[19] == 1
    assert sum(entries) == 6
    assert exits[9] == 2 and exits[17] == 1 and exits[19] == 2
    assert sum(exits) == 5


def test_summary(filled):
    report = ReportService(filled).summary(DAY, DAY)
    assert report.entries == 6
    assert report.exits == 5
    assert report.unique_plates == 6
    assert report.revenue == Decimal("375.00")
    assert report.unpaid_count == 1
    assert report.unpaid_amount == Decimal("90.00")
    assert report.by_type[VehicleType.CAR] == 3
    assert report.by_type[VehicleType.UNKNOWN] == 1
    assert report.revenue_by_method[PaymentMethod.CARD] == Decimal("315.00")
    assert report.revenue_by_method[PaymentMethod.CASH] == Decimal("60.00")
    assert report.peak_hour in (8, 17)
    assert report.average_stay == timedelta(minutes=(90 + 60 + 120 + 10 + 61) / 5)
    assert [d.entries for d in report.days] == [6]
    assert report.average_daily_revenue == Decimal("375.00")


def test_summary_swaps_reversed_range_and_handles_empty(db):
    report = ReportService(db).summary(date(2026, 1, 10), date(2026, 1, 1))
    assert report.start == date(2026, 1, 1)
    assert report.day_count == 10
    assert report.entries == 0
    assert report.peak_hour is None


def test_exports(filled, tmp_path):
    sessions = sessions_for_export(filled, DAY, DAY)
    assert len(sessions) == 6

    csv_path = export_sessions_csv(sessions, tmp_path / "gecisler.csv")
    raw = csv_path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # Excel için BOM
    rows = list(csv.reader(csv_path.open(encoding="utf-8-sig"), delimiter=";"))
    assert rows[0][1] == "Plaka"
    assert rows[1][1] == "34 A 1111"
    assert rows[1][8] == "90,00"

    summary = ReportService(filled).summary(DAY, DAY)
    xlsx = export_sessions_xlsx(sessions, tmp_path / "rapor.xlsx", summary=summary)
    workbook = load_workbook(xlsx)
    assert workbook.sheetnames == ["Geçişler", "Günlük Özet", "Özet"]
    sheet = workbook["Geçişler"]
    assert sheet.max_row == 7
    assert sheet["B2"].value == "34 A 1111"
    assert sheet["I2"].value == 90
    assert isinstance(sheet["E2"].value, datetime)
    assert workbook["Günlük Özet"]["D2"].value == 375


def test_password_hashing():
    stored = hash_password("gizli-1234")
    assert stored.startswith("scrypt$")
    assert verify_password("gizli-1234", stored)
    assert not verify_password("yanlis", stored)
    assert not verify_password("x", "bozuk")
    assert hash_password("a1b2") != hash_password("a1b2")  # tuzlu


def test_user_service(db):
    users = UserService(db)
    assert users.ensure_default_admin()
    assert not users.ensure_default_admin()
    admin = users.authenticate("Admin", "admin")
    assert admin is not None and admin.must_change_password and admin.last_login
    assert users.authenticate("admin", "yanlis") is None

    users.change_password(admin.id, "admin", "yeni-parola")
    assert users.authenticate("admin", "yeni-parola").must_change_password is False
    with pytest.raises(AuthError):
        users.change_password(admin.id, "admin", "x1234")

    op = users.create("ali.veli", "1234", full_name="Ali Veli")
    assert op.role is UserRole.OPERATOR
    with pytest.raises(AuthError):
        users.create("ali.veli", "1234")
    with pytest.raises(AuthError):
        users.create("kötü isim!", "1234")
    with pytest.raises(AuthError):
        users.create("kisa", "12")

    # Son yönetici silinemez / yetkisi düşürülemez / pasifleştirilemez
    with pytest.raises(AuthError):
        users.delete(admin.id)
    with pytest.raises(AuthError):
        users.update(admin.id, role=UserRole.OPERATOR)
    with pytest.raises(AuthError):
        users.update(admin.id, is_active=False)

    users.update(op.id, role=UserRole.ADMIN)
    users.update(admin.id, is_active=False)
    assert users.authenticate("admin", "yeni-parola") is None
    users.set_password(op.id, "sifirla", must_change=True)
    assert users.authenticate("ali.veli", "sifirla").must_change_password
    users.delete(admin.id)
    assert [u.username for u in users.list()] == ["ali.veli"]


def test_backup(tmp_path):
    db = Database(f"sqlite:///{tmp_path / 'otopark.db'}")
    db.create_schema()
    seed_demo_data(db, days=1, rng=random.Random(1))
    backups = tmp_path / "yedek"
    path = backup_database(db, backups, keep=2)
    assert path.is_file()
    copy = Database(f"sqlite:///{path}")
    assert len(ParkingService(copy, AppConfig()).find_sessions()) > 0
    copy.dispose()
    assert backup_if_due(db, backups) is None  # bugün zaten yedeklendi
    assert len(list_backups(backups)) == 1
    with pytest.raises(BackupError):
        backup_database(Database("sqlite://"), backups)
    db.dispose()
